import re
import zipfile
from pathlib import Path

p = Path(r"c:\flutter_projects\organizzazione_personale_cronos - mod\assets\Mod.RFP_03.xlsx")
with zipfile.ZipFile(p) as z:
    print("MEDIA:")
    for n in sorted(z.namelist()):
        if "media/" in n:
            print(f"  {n} {z.getinfo(n).file_size}")

    sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    merges = re.findall(r'mergeCell ref="([^"]+)"', sheet)
    print(f"\nMERGES ({len(merges)}):")
    for m in merges:
        print(f"  {m}")

    cols = re.findall(r'<col min="(\d+)" max="(\d+)" width="([^"]+)"', sheet)
    print("\nCOL WIDTHS:")
    for mn, mx, w in cols[:20]:
        print(f"  {mn}-{mx}: {w}")

    rows = re.findall(r'<row r="(\d+)"([^>]*)>', sheet)
    print("\nROWS with ht/customHeight:")
    for r, attrs in rows:
        if "ht=" in attrs or "customHeight" in attrs:
            ht = re.search(r'ht="([^"]+)"', attrs)
            print(f"  row {r} ht={ht.group(1) if ht else '?'}")

    # cell values rows 11-55
    try:
        import openpyxl
        wb = openpyxl.load_workbook(p)
        ws = wb.active
        print("\nCELL VALUES rows 11-55:")
        for r in range(11, 56):
            parts = []
            for c in range(1, 20):
                from openpyxl.utils import get_column_letter
                ref = f"{get_column_letter(c)}{r}"
                v = ws[ref].value
                if v is not None and str(v).strip():
                    s = str(v).replace("\n", " ")[:50]
                    parts.append(f"{ref}={s}")
            if parts:
                print(f"R{r}: " + " | ".join(parts))
    except ImportError:
        print("openpyxl not installed")

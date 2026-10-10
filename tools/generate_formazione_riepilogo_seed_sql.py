from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


SRC = Path(r"c:\Users\alexandru.cibuc\Downloads\CORSI DI FORMAZIONE.xlsx")
OUT = Path(
    r"c:\flutter_projects\organizzazione_personale_cronos - mod\supabase\migrations\20260414173000_seed_formazione_corsi_from_riepilogo.sql"
)


def to_iso(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return ""
    s = str(v).strip()
    if not s:
        return ""
    try:
        return datetime.fromisoformat(s.replace("Z", "")).date().isoformat()
    except Exception:
        return ""


def s(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def q(v: str) -> str:
    return "'" + v.replace("'", "''") + "'"


def main() -> None:
    wb = load_workbook(SRC, data_only=True)
    ws = wb["RIEPILOGO"]

    # Mappa corsi: riga 1 nome corso, riga 2 ATT/SCAD
    groups: list[tuple[str, int, int | None]] = []
    col = 5
    while col <= ws.max_column:
        g = s(ws.cell(1, col).value)
        if not g:
            col += 1
            continue
        att_col = col
        scad_col = col + 1 if s(ws.cell(2, col + 1).value).upper().startswith("SCAD") else None
        groups.append((g, att_col, scad_col))
        col += 2 if scad_col else 1

    rows = []
    for r in range(3, ws.max_row + 1):
        nome = s(ws.cell(r, 4).value)
        if not nome:
            continue
        dimessi = s(ws.cell(r, 3).value)
        for corso, att_c, scad_c in groups:
            att_raw = ws.cell(r, att_c).value
            scad_raw = ws.cell(r, scad_c).value if scad_c else None
            att = s(att_raw)
            scad = to_iso(scad_raw)
            att_date = to_iso(att_raw)
            # prende solo celle valorizzate
            if not att and not scad and not att_date:
                continue
            rows.append(
                {
                    "nome": nome,
                    "corso": corso,
                    "attestato": att if not att_date else "",
                    "data_attestato": att_date,
                    "scadenza_attestato": scad,
                    "dimessi": dimessi,
                }
            )

    values = []
    for x in rows:
        values.append(
            "("
            + ",".join(
                [
                    q(x["nome"]),
                    q(x["corso"]),
                    q(x["attestato"]),
                    q(x["data_attestato"]) if x["data_attestato"] else "null",
                    q(x["scadenza_attestato"]) if x["scadenza_attestato"] else "null",
                    q(x["dimessi"]),
                ]
            )
            + ")"
        )

    sql = """-- Seed formazione_corsi da foglio RIEPILOGO (copertura nominativi completa)
with src(nome, corso, attestato, data_attestato, scadenza_attestato, dimessi) as (
  values
"""
    sql += ",\n".join(values)
    sql += """
), normalized as (
  select
    trim(regexp_replace(lower(nome), '\\s+', ' ', 'g')) as nome_norm,
    trim(regexp_replace(lower(corso), '\\s+', ' ', 'g')) as corso_norm,
    *
  from src
), personale_match as (
  select
    p.id_uuid as personale_id,
    trim(regexp_replace(lower(p.full_name), '\\s+', ' ', 'g')) as nome_norm
  from public.personale p
), mapped as (
  select
    pm.personale_id,
    n.corso,
    nullif(n.attestato, '') as attestato,
    n.data_attestato::date as data_attestato,
    n.scadenza_attestato::date as scadenza_attestato,
    nullif(n.dimessi, '') as dimessi
  from normalized n
  join personale_match pm on pm.nome_norm = n.nome_norm
), dedup as (
  select distinct on (personale_id, corso, coalesce(data_attestato, '1900-01-01'::date), coalesce(scadenza_attestato, '1900-01-01'::date))
    *
  from mapped
  order by personale_id, corso, coalesce(data_attestato, '1900-01-01'::date), coalesce(scadenza_attestato, '1900-01-01'::date)
)
insert into public.formazione_corsi (
  personale_id, corso, attestato, data_attestato, scadenza_attestato, dimessi
)
select personale_id, corso, attestato, data_attestato, scadenza_attestato, dimessi
from dedup
on conflict (personale_id, corso, data_attestato) do update set
  attestato = excluded.attestato,
  scadenza_attestato = excluded.scadenza_attestato,
  dimessi = excluded.dimessi,
  updated_at = now();
"""

    OUT.write_text(sql, encoding="utf-8")
    print(f"Written {OUT} with {len(rows)} row-cells")


if __name__ == "__main__":
    main()


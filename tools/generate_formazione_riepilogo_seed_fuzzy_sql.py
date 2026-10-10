import difflib
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

import requests
from openpyxl import load_workbook

SUPABASE_URL = "https://bjdimalvbdablzoctrsf.supabase.co"
ANON_KEY = (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImJqZGltYWx2YmRhYmx6b2N0cnNmIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzE5MzczNTcsImV4cCI6MjA4NzUxMzM1N30."
    "Lm4GHd1CWdSMMI87Wx2eBrHcmqFtMXO31hO7Q3BsYyc"
)
SRC = Path(r"c:\Users\alexandru.cibuc\Downloads\CORSI DI FORMAZIONE.xlsx")
OUT = Path(
    r"c:\flutter_projects\organizzazione_personale_cronos - mod\supabase\migrations\20260414175000_seed_formazione_corsi_riepilogo_fuzzy.sql"
)


def norm(text: str) -> str:
    t = (text or "").strip().lower().replace("’", "'")
    t = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")
    t = re.sub(r"[^a-z0-9 ]+", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def token_key(text: str) -> str:
    parts = [p for p in norm(text).split(" ") if p]
    parts.sort()
    return " ".join(parts)


def to_iso(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return ""
    s = str(v).strip()
    if not s:
        return ""
    try:
        return datetime.fromisoformat(s.replace("Z", "")).date().isoformat()
    except Exception:
        return ""


def s(v) -> str:
    return "" if v is None else str(v).strip()


def q(v: str) -> str:
    return "'" + v.replace("'", "''") + "'"


def best_match(name: str, people: list[dict]) -> str | None:
    n1 = norm(name)
    n2 = token_key(name)
    best_id = None
    best_score = 0.0
    for p in people:
        sc1 = difflib.SequenceMatcher(None, n1, p["norm"]).ratio()
        sc2 = difflib.SequenceMatcher(None, n2, p["token"]).ratio()
        score = max(sc1, sc2)
        if score > best_score:
            best_score = score
            best_id = p["id"]
    return best_id if best_score >= 0.90 else None


def main() -> None:
    headers = {"apikey": ANON_KEY, "Authorization": f"Bearer {ANON_KEY}"}
    resp = requests.get(
        f"{SUPABASE_URL}/rest/v1/personale?select=id_uuid,full_name&limit=5000",
        headers=headers,
        timeout=30,
    )
    resp.raise_for_status()
    pers_raw = resp.json()
    people = [
        {
            "id": (p.get("id_uuid") or "").strip(),
            "norm": norm(p.get("full_name") or ""),
            "token": token_key(p.get("full_name") or ""),
        }
        for p in pers_raw
        if (p.get("id_uuid") or "").strip() and (p.get("full_name") or "").strip()
    ]
    exact_map = {p["norm"]: p["id"] for p in people}
    token_map = {p["token"]: p["id"] for p in people}

    wb = load_workbook(SRC, data_only=True)
    ws = wb["RIEPILOGO"]

    groups = []
    c = 5
    while c <= ws.max_column:
        g = s(ws.cell(1, c).value)
        if not g:
            c += 1
            continue
        scad_col = c + 1 if s(ws.cell(2, c + 1).value).upper().startswith("SCAD") else None
        groups.append((g, c, scad_col))
        c += 2 if scad_col else 1

    rows = []
    unmatched = set()
    for r in range(3, ws.max_row + 1):
        nome = s(ws.cell(r, 4).value)
        if not nome:
            continue
        pid = exact_map.get(norm(nome)) or token_map.get(token_key(nome)) or best_match(nome, people)
        if not pid:
            unmatched.add(nome)
            continue
        dimessi = s(ws.cell(r, 3).value)
        for corso, att_c, scad_c in groups:
            att_raw = ws.cell(r, att_c).value
            scad_raw = ws.cell(r, scad_c).value if scad_c else None
            att = s(att_raw)
            att_date = to_iso(att_raw)
            scad = to_iso(scad_raw)
            if not att and not att_date and not scad:
                continue
            rows.append(
                {
                    "pid": pid,
                    "corso": corso,
                    "att": "" if att_date else att,
                    "att_date": att_date,
                    "scad": scad,
                    "dimessi": dimessi,
                }
            )

    values = []
    for x in rows:
        values.append(
            "("
            + ",".join(
                [
                    q(x["pid"]),
                    q(x["corso"]),
                    q(x["att"]),
                    q(x["att_date"]) if x["att_date"] else "null",
                    q(x["scad"]) if x["scad"] else "null",
                    q(x["dimessi"]),
                ]
            )
            + ")"
        )

    sql = """-- Seed fuzzy da RIEPILOGO per aumentare copertura nomi
with src(personale_id, corso, attestato, data_attestato, scadenza_attestato, dimessi) as (
  values
"""
    sql += ",\n".join(values)
    sql += """
), dedup as (
  select distinct on (
    personale_id,
    corso,
    coalesce(data_attestato::date, '1900-01-01'::date),
    coalesce(scadenza_attestato::date, '1900-01-01'::date)
  )
    personale_id::uuid, corso, nullif(attestato,''), data_attestato::date, scadenza_attestato::date, nullif(dimessi,'')
  from src
  order by
    personale_id,
    corso,
    coalesce(data_attestato::date, '1900-01-01'::date),
    coalesce(scadenza_attestato::date, '1900-01-01'::date)
)
insert into public.formazione_corsi (
  personale_id, corso, attestato, data_attestato, scadenza_attestato, dimessi
)
select * from dedup
on conflict (personale_id, corso, data_attestato) do update set
  attestato = excluded.attestato,
  scadenza_attestato = excluded.scadenza_attestato,
  dimessi = excluded.dimessi,
  updated_at = now();
"""

    OUT.write_text(sql, encoding="utf-8")
    print("people:", len(people))
    print("rows:", len(rows))
    print("unmatched:", len(unmatched))
    print("unmatched sample:", sorted(unmatched)[:30])
    print("written:", OUT)


if __name__ == "__main__":
    main()


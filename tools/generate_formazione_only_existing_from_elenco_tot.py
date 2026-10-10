from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(r"c:\Users\alexandru.cibuc\Downloads\CORSI DI FORMAZIONE.xlsx")
OUT = Path(
    r"c:\flutter_projects\organizzazione_personale_cronos - mod\supabase\migrations\20260414195500_sync_formazione_only_existing_names.sql"
)


def s(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def d(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    t = s(v)
    if not t:
        return ""
    try:
        return datetime.fromisoformat(t.replace("Z", "")).date().isoformat()
    except Exception:
        return ""


def q(v: str) -> str:
    return "'" + v.replace("'", "''") + "'"


def main() -> None:
    wb = load_workbook(SRC, data_only=True)
    ws = wb["ELENCO TOT"]

    rows = []
    for r in range(2, ws.max_row + 1):
        nome = s(ws.cell(r, 7).value)  # DIPENDENTE
        corso = s(ws.cell(r, 5).value)  # CORSO
        if not nome or not corso:
            continue
        anno = s(ws.cell(r, 3).value)
        rows.append(
            {
                "nome": nome,
                "anno": anno if anno.isdigit() else "",
                "data": d(ws.cell(r, 1).value),
                "oda": s(ws.cell(r, 2).value),
                "ente": s(ws.cell(r, 4).value),
                "corso": corso,
                "primo": s(ws.cell(r, 6).value),
                "attestato": s(ws.cell(r, 10).value),
                "data_att": d(ws.cell(r, 11).value),
                "scadenza": d(ws.cell(r, 12).value),
                "dimessi": s(ws.cell(r, 13).value),
                "note": s(ws.cell(r, 14).value),
            }
        )

    values = []
    for x in rows:
        values.append(
            "("
            + ",".join(
                [
                    q(x["nome"]),
                    x["anno"] if x["anno"] else "null",
                    q(x["data"]) if x["data"] else "null",
                    q(x["oda"]),
                    q(x["ente"]),
                    q(x["corso"]),
                    q(x["primo"]),
                    q(x["attestato"]),
                    q(x["data_att"]) if x["data_att"] else "null",
                    q(x["scadenza"]) if x["scadenza"] else "null",
                    q(x["dimessi"]),
                    q(x["note"]),
                ]
            )
            + ")"
        )

    sql = """-- Allinea solo corsi formazione per nomi gia presenti in personale
-- NON inserisce nuovi nominativi.
create extension if not exists unaccent;

with src(
  nome, anno, data, oda, ente, corso, primo_rilascio_aggiornamento, attestato,
  data_attestato, scadenza_attestato, dimessi, note
) as (
  values
"""
    sql += ",\n".join(values)
    sql += """
), normalized_src as (
  select
    (
      select string_agg(tok, ' ' order by tok)
      from unnest(
        regexp_split_to_array(
          regexp_replace(
            lower(
              unaccent(
                replace(replace(replace(coalesce(nome, ''), '’', ''''), '´', ''''), '`', '''')
              )
            ),
            '[^a-z0-9 ]+',
            ' ',
            'g'
          ),
          '\s+'
        )
      ) as tok
      where tok <> ''
    ) as token_key,
    *
  from src
), normalized_personale as (
  select
    p.id_uuid as personale_id,
    (
      select string_agg(tok, ' ' order by tok)
      from unnest(
        regexp_split_to_array(
          regexp_replace(
            lower(
              unaccent(
                replace(replace(replace(coalesce(p.full_name, ''), '’', ''''), '´', ''''), '`', '''')
              )
            ),
            '[^a-z0-9 ]+',
            ' ',
            'g'
          ),
          '\s+'
        )
      ) as tok
      where tok <> ''
    ) as token_key
  from public.personale p
), mapped as (
  select
    np.personale_id,
    ns.anno::integer,
    ns.data::date,
    nullif(ns.oda, '') as oda,
    nullif(ns.ente, '') as ente,
    ns.corso,
    nullif(ns.primo_rilascio_aggiornamento, '') as primo_rilascio_aggiornamento,
    nullif(ns.attestato, '') as attestato,
    ns.data_attestato::date as data_attestato,
    ns.scadenza_attestato::date as scadenza_attestato,
    nullif(ns.dimessi, '') as dimessi,
    nullif(ns.note, '') as note
  from normalized_src ns
  join normalized_personale np on np.token_key = ns.token_key
  where coalesce(ns.token_key, '') <> ''
), dedup as (
  select distinct on (
    personale_id,
    corso,
    coalesce(data_attestato, '1900-01-01'::date),
    coalesce(scadenza_attestato, '1900-01-01'::date)
  )
    personale_id,
    anno,
    data,
    oda,
    ente,
    corso,
    primo_rilascio_aggiornamento,
    attestato,
    data_attestato,
    scadenza_attestato,
    null::text as corso_prenotato,
    null::date as prima_data,
    null::date as seconda_data,
    null::text as orario,
    null::text as modalita,
    dimessi,
    note
  from mapped
  order by
    personale_id,
    corso,
    coalesce(data_attestato, '1900-01-01'::date),
    coalesce(scadenza_attestato, '1900-01-01'::date),
    anno desc nulls last,
    data desc nulls last
)
insert into public.formazione_corsi (
  personale_id, anno, data, oda, ente, corso, primo_rilascio_aggiornamento, attestato,
  data_attestato, scadenza_attestato, corso_prenotato, prima_data, seconda_data, orario, modalita, dimessi, note
)
select * from dedup
on conflict (personale_id, corso, data_attestato) do update set
  anno = excluded.anno,
  data = excluded.data,
  oda = excluded.oda,
  ente = excluded.ente,
  primo_rilascio_aggiornamento = excluded.primo_rilascio_aggiornamento,
  attestato = excluded.attestato,
  scadenza_attestato = excluded.scadenza_attestato,
  dimessi = excluded.dimessi,
  note = excluded.note,
  updated_at = now();
"""

    OUT.write_text(sql, encoding="utf-8")
    print(f"Rows from ELENCO TOT: {len(rows)}")
    print(f"Written migration: {OUT}")


if __name__ == "__main__":
    main()


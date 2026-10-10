import re
import unicodedata
from pathlib import Path

import requests
from openpyxl import load_workbook


SUPABASE_URL = "https://bjdimalvbdablzoctrsf.supabase.co"
ANON_KEY = (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImJqZGltYWx2YmRhYmx6b2N0cnNmIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzE5MzczNTcsImV4cCI6MjA4NzUxMzM1N30."
    "Lm4GHd1CWdSMMI87Wx2eBrHcmqFtMXO31hO7Q3BsYyc"
)
XLSX = Path(r"c:\Users\alexandru.cibuc\Downloads\CORSI DI FORMAZIONE.xlsx")


def norm(text: str) -> str:
    t = (text or "").strip().lower().replace("’", "'")
    t = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")
    t = re.sub(r"[^a-z0-9 ]+", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def token_key(text: str) -> str:
    tokens = [x for x in norm(text).split(" ") if x]
    tokens.sort()
    return " ".join(tokens)


def main() -> None:
    headers = {"apikey": ANON_KEY, "Authorization": f"Bearer {ANON_KEY}"}
    r = requests.get(f"{SUPABASE_URL}/rest/v1/personale?select=id_uuid,full_name&limit=5000", headers=headers, timeout=30)
    r.raise_for_status()
    personale = r.json()
    print("personale rows:", len(personale))

    p_exact = {norm(p.get("full_name", "")) for p in personale if p.get("full_name")}
    p_token = {token_key(p.get("full_name", "")) for p in personale if p.get("full_name")}

    wb = load_workbook(XLSX, data_only=True)
    ws = wb["RIEPILOGO"]
    excel_names = []
    for row in range(3, ws.max_row + 1):
        nome = ws.cell(row, 4).value
        if nome and str(nome).strip():
            excel_names.append(str(nome).strip())
    uniq = sorted(set(excel_names))
    print("excel unique names:", len(uniq))

    exact_ok = []
    token_ok = []
    miss = []
    for n in uniq:
        nn = norm(n)
        tk = token_key(n)
        if nn in p_exact:
            exact_ok.append(n)
        elif tk in p_token:
            token_ok.append(n)
        else:
            miss.append(n)

    print("exact matches:", len(exact_ok))
    print("token matches:", len(token_ok))
    print("unmatched:", len(miss))
    print("unmatched sample:", miss[:40])


if __name__ == "__main__":
    main()


import zipfile
import re
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

p = Path(r"c:\flutter_projects\organizzazione_personale_cronos - mod\assets\Mod.RFP_03.xlsx")
wb = openpyxl.load_workbook(p)
ws = wb.active

print("Row 11 fill:", ws["B11"].fill.fgColor.rgb if ws["B11"].fill else None)
print("Row 13 fill:", ws["B13"].fill.fgColor.rgb if ws["B13"].fill else None)
print("Row 16 B fill:", ws["B16"].fill.fgColor.rgb if ws["B16"].fill else None)

# footer text rows
for r in range(52, 62):
    parts = []
    for c in range(2, 15):
        ref = f"{get_column_letter(c)}{r}"
        v = ws[ref].value
        if v:
            parts.append(str(v)[:60])
    if parts:
        print(f"R{r}:", " | ".join(parts))

with zipfile.ZipFile(p) as z:
    theme = z.read("xl/theme/theme1.xml").decode("utf-8")
    for i, m in enumerate(re.finditer(r'<a:srgbClr val="([A-F0-9]+)"', theme)):
        print(f"theme color {i}: #{m.group(1)}")

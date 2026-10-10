from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


SRC = Path(r"c:\Users\alexandru.cibuc\Downloads\CORSI DI FORMAZIONE.xlsx")
OUT = Path(
    r"c:\flutter_projects\organizzazione_personale_cronos - mod\supabase\migrations\20260414170000_seed_formazione_corsi_from_excel.sql"
)


def date_cell(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return ""
    t = str(v).strip()
    if not t:
        return ""
    try:
        return datetime.fromisoformat(t.replace("Z", "")).date().isoformat()
    except Exception:
        return ""


def s_cell(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def sql_str(v: str) -> str:
    return "'" + v.replace("'", "''") + "'"


def main() -> None:
    wb = load_workbook(SRC, data_only=True)
    rows: list[dict[str, str]] = []
    for sheet in wb.sheetnames:
        if not (len(sheet) == 4 and sheet.isdigit()):
            continue
        ws = wb[sheet]
        anno = int(sheet)
        for r in range(3, ws.max_row + 1):
            nome = s_cell(ws.cell(r, 7).value)
            corso = s_cell(ws.cell(r, 5).value)
            if not nome or not corso:
                continue
            rows.append(
                {
                    "nome": nome,
                    "anno": str(anno),
                    "data": date_cell(ws.cell(r, 1).value),
                    "oda": s_cell(ws.cell(r, 2).value),
                    "ente": s_cell(ws.cell(r, 4).value),
                    "corso": corso,
                    "primo": s_cell(ws.cell(r, 6).value),
                    "attestato": s_cell(ws.cell(r, 8).value),
                    "data_att": date_cell(ws.cell(r, 9).value),
                    "scadenza": date_cell(ws.cell(r, 10).value),
                    "prenotato": s_cell(ws.cell(r, 11).value),
                    "prima_data": date_cell(ws.cell(r, 12).value),
                    "seconda_data": date_cell(ws.cell(r, 13).value),
                    "orario": s_cell(ws.cell(r, 14).value),
                    "modalita": s_cell(ws.cell(r, 15).value),
                    "dimessi": s_cell(ws.cell(r, 16).value),
                    "note": s_cell(ws.cell(r, 17).value),
                }
            )

    values_sql = []
    for x in rows:
        values_sql.append(
            "("
            + ",".join(
                [
                    sql_str(x["nome"]),
                    x["anno"],
                    sql_str(x["data"]) if x["data"] else "null",
                    sql_str(x["oda"]),
                    sql_str(x["ente"]),
                    sql_str(x["corso"]),
                    sql_str(x["primo"]),
                    sql_str(x["attestato"]),
                    sql_str(x["data_att"]) if x["data_att"] else "null",
                    sql_str(x["scadenza"]) if x["scadenza"] else "null",
                    sql_str(x["prenotato"]),
                    sql_str(x["prima_data"]) if x["prima_data"] else "null",
                    sql_str(x["seconda_data"]) if x["seconda_data"] else "null",
                    sql_str(x["orario"]),
                    sql_str(x["modalita"]),
                    sql_str(x["dimessi"]),
                    sql_str(x["note"]),
                ]
            )
            + ")"
        )

    sql = """-- Seed formazione_corsi da file CORSI DI FORMAZIONE.xlsx
with src(
  nome, anno, data, oda, ente, corso, primo_rilascio_aggiornamento, attestato,
  data_attestato, scadenza_attestato, corso_prenotato, prima_data, seconda_data, orario, modalita, dimessi, note
) as (
  values
"""
    sql += ",\n".join(values_sql)
    sql += """
), normalized as (
  select
    trim(regexp_replace(lower(nome), '\\s+', ' ', 'g')) as nome_norm,
    *
  from src
), personale_match as (
  select
    p.id_uuid as personale_id,
    trim(regexp_replace(lower(p.full_name), '\\s+', ' ', 'g')) as nome_norm
  from public.personale p
), mapped as (
  select
    pm.personale_id,
    n.anno::integer,
    n.data::date,
    nullif(n.oda, ''),
    nullif(n.ente, ''),
    n.corso,
    nullif(n.primo_rilascio_aggiornamento, ''),
    nullif(n.attestato, ''),
    n.data_attestato::date,
    n.scadenza_attestato::date,
    nullif(n.corso_prenotato, ''),
    n.prima_data::date,
    n.seconda_data::date,
    nullif(n.orario, ''),
    nullif(n.modalita, ''),
    nullif(n.dimessi, ''),
    nullif(n.note, '')
  from normalized n
  join personale_match pm on pm.nome_norm = n.nome_norm
), dedup as (
  select distinct on (personale_id, corso, data_attestato)
    *
  from mapped
  order by personale_id, corso, data_attestato, anno desc, data desc nulls last
)
insert into public.formazione_corsi (
  personale_id, anno, data, oda, ente, corso, primo_rilascio_aggiornamento, attestato,
  data_attestato, scadenza_attestato, corso_prenotato, prima_data, seconda_data, orario, modalita, dimessi, note
)
select * from dedup
on conflict (personale_id, corso, data_attestato) do update set
  anno = excluded.anno,
  data = excluded.data,
  oda = excluded.oda,
  ente = excluded.ente,
  primo_rilascio_aggiornamento = excluded.primo_rilascio_aggiornamento,
  attestato = excluded.attestato,
  scadenza_attestato = excluded.scadenza_attestato,
  corso_prenotato = excluded.corso_prenotato,
  prima_data = excluded.prima_data,
  seconda_data = excluded.seconda_data,
  orario = excluded.orario,
  modalita = excluded.modalita,
  dimessi = excluded.dimessi,
  note = excluded.note,
  updated_at = now();
"""

    OUT.write_text(sql, encoding="utf-8")
    print(f"Written {OUT} with {len(rows)} rows")


if __name__ == "__main__":
    main()


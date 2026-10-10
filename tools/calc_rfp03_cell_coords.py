"""Calcola coordinate PDF (pt) per overlay celle Mod.RFP_03."""
import json
import re
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

XLSX = r"c:\flutter_projects\organizzazione_personale_cronos - mod\assets\Mod.RFP_03.xlsx"
PAGE_W = 595.28
PAGE_H = 841.89

wb = openpyxl.load_workbook(XLSX)
ws = wb.active
m = ws.page_margins
ml = (m.left or 0) * 72
mr = (m.right or 0) * 72
mt = (m.top or 0) * 72
mb = (m.bottom or 0) * 72

PRINT_COLS = list(range(1, 21))  # A-T
PRINT_ROWS = list(range(1, 65))

col_pt = []
for c in PRINT_COLS:
    letter = get_column_letter(c)
    w = ws.column_dimensions[letter].width or 8.43
    col_pt.append(max(1.0, w * 7.0 + 5.0) * 72.0 / 96.0)

row_pt = []
for r in PRINT_ROWS:
    row_pt.append(ws.row_dimensions[r].height or 15.0)

content_w = sum(col_pt)
content_h = sum(row_pt)
print_w = PAGE_W - ml - mr
print_h = PAGE_H - mt - mb
scale = min(print_w / content_w, print_h / content_h)
sw = [w * scale for w in col_pt]
sh = [h * scale for h in row_pt]
offset_x = ml + (print_w - sum(sw)) / 2
offset_y = PAGE_H - mt - (print_h - sum(sh)) / 2

col_x = [offset_x]
for w in sw:
    col_x.append(col_x[-1] + w)

row_y_top = [offset_y]
for h in sh:
    row_y_top.append(row_y_top[-1] - h)


def merged_rect(ref: str):
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    x0 = col_x[min_col - 1]
    x1 = col_x[max_col]
    y_top = row_y_top[min_row - 1]
    y_bot = row_y_top[max_row]
    return {
        "x": round(x0, 2),
        "y": round(y_bot, 2),
        "w": round(x1 - x0, 2),
        "h": round(y_top - y_bot, 2),
    }


OVERLAYS = {
    "B13": "B13:S14",
    "B16": "B16:E17",
    "F16": "F16:J17",
    "K16": "K16:N17",
    "O16": "O16:S17",
    "K18": "K18:K19",
    "N18": "N18:N19",
    "R18": "R18:S19",
    "D37": "D37",
    "D38": "D38",
    "D40": "D40",
    "E40": "E40:S40",
    "C44": "C44:H44",
    "D49": "D49:G49",
    "N46": "N46:Q46",
    "E46": "D46:G49",
}

out = {"page": {"w": PAGE_W, "h": PAGE_H}, "cells": {}}
for key, merge in OVERLAYS.items():
    out["cells"][key] = merged_rect(merge)

OUT = Path(__file__).resolve().parents[1] / "supabase/functions/render-rfp03-pdf/cell_coords.json"
OUT.write_text(json.dumps(out, indent=2), encoding="utf-8")
print(f"Wrote {OUT}")
print(json.dumps(out, indent=2))

try:
    import pdfplumber
    with pdfplumber.open(
        r"c:\flutter_projects\organizzazione_personale_cronos - mod\assets\Mod.RFP_03_background.pdf"
    ) as pdf:
        words = pdf.pages[0].extract_words()
        for t in ["DATA:", "Richiesta", "Di:", "Patrucco"]:
            hits = [w for w in words if t in w["text"]]
            if hits:
                w = hits[0]
                print(f"CALIB {t}: x0={w['x0']:.1f} top={w['top']:.1f}", file=__import__("sys").stderr)
except ImportError:
    pass

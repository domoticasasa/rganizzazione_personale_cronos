import re
import unicodedata

import requests
from openpyxl import load_workbook

SUPABASE_URL = "https://bjdimalvbdablzoctrsf.supabase.co"
ANON_KEY = (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImJqZGltYWx2YmRhYmx6b2N0cnNmIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzE5MzczNTcsImV4cCI6MjA4NzUxMzM1N30."
    "Lm4GHd1CWdSMMI87Wx2eBrHcmqFtMXO31hO7Q3BsYyc"
)
XLSX = r"c:\Users\alexandru.cibuc\Downloads\CORSI DI FORMAZIONE.xlsx"


def tkey(text: str) -> str:
    s = (text or "").strip().lower().replace("’", "'")
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    parts = [p for p in s.split(" ") if p]
    parts.sort()
    return " ".join(parts)


def main() -> None:
    h = {"apikey": ANON_KEY, "Authorization": f"Bearer {ANON_KEY}"}

    # Discover likely users/profile table
    candidates = [
        "users",
        "utenti",
        "user_profiles",
        "profiles",
        "app_users",
    ]
    found = None
    for t in candidates:
        r = requests.get(
            f"{SUPABASE_URL}/rest/v1/{t}?select=*&limit=1",
            headers=h,
            timeout=20,
        )
        if r.status_code == 200:
            found = t
            print("admin table candidate:", t)
            print("sample row keys:", list((r.json()[0] if r.json() else {}).keys()))
            break
    if not found:
        print("admin table candidate: NOT FOUND among", candidates)

    p = requests.get(
        f"{SUPABASE_URL}/rest/v1/personale?select=id_uuid,full_name,active,email,user_id&limit=5000",
        headers=h,
        timeout=30,
    )
    p.raise_for_status()
    personale = p.json()
    print("personale rows:", len(personale))

    wb = load_workbook(XLSX, data_only=True)
    ws = wb["RIEPILOGO"]
    excel_names = {
        str(ws.cell(r, 4).value).strip()
        for r in range(3, ws.max_row + 1)
        if ws.cell(r, 4).value
    }
    print("riepilogo names:", len(excel_names))

    pkeys = {tkey((x.get("full_name") or "")) for x in personale}
    missing = [n for n in sorted(excel_names) if tkey(n) not in pkeys]
    print("missing vs personale:", len(missing))
    print("missing sample:", missing[:40])


if __name__ == "__main__":
    main()


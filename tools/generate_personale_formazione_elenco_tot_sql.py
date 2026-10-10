from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(r"c:\Users\alexandru.cibuc\Downloads\CORSI DI FORMAZIONE.xlsx")
OUT = Path(
    r"c:\flutter_projects\organizzazione_personale_cronos - mod\supabase\migrations\20260414182000_sync_personale_173_and_seed_formazione_elenco_tot.sql"
)


def s(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def d(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    t = s(v)
    if not t:
        return ""
    try:
        return datetime.fromisoformat(t.replace("Z", "")).date().isoformat()
    except Exception:
        return ""


def q(v: str) -> str:
    return "'" + v.replace("'", "''") + "'"


def main() -> None:
    wb = load_workbook(SRC, data_only=True)

    # 1) Nominativi da RIEPILOGO (173)
    ws_riep = wb["RIEPILOGO"]
    personale_names = sorted(
        {
            s(ws_riep.cell(r, 4).value)
            for r in range(3, ws_riep.max_row + 1)
            if s(ws_riep.cell(r, 4).value)
        }
    )

    # 2) Righe corsi da ELENCO TOT (completo per ente + data attestato)
    ws_tot = wb["ELENCO TOT"]
    formazione_rows: list[dict[str, str]] = []
    for r in range(2, ws_tot.max_row + 1):
        nome = s(ws_tot.cell(r, 7).value)  # DIPENDENTE
        corso = s(ws_tot.cell(r, 5).value)  # CORSO
        if not nome or not corso:
            continue
        formazione_rows.append(
            {
                "nome": nome,
                "anno": s(ws_tot.cell(r, 3).value),
                "data": d(ws_tot.cell(r, 1).value),
                "oda": s(ws_tot.cell(r, 2).value),
                "ente": s(ws_tot.cell(r, 4).value),
                "corso": corso,
                "primo": s(ws_tot.cell(r, 6).value),
                "attestato": s(ws_tot.cell(r, 10).value),
                "data_att": d(ws_tot.cell(r, 11).value),
                "scadenza": d(ws_tot.cell(r, 12).value),
                "dimessi": s(ws_tot.cell(r, 13).value),
                "note": s(ws_tot.cell(r, 14).value),
            }
        )

    personale_values = ",\n".join(f"({q(n)})" for n in personale_names)
    formazione_values = ",\n".join(
        "("
        + ",".join(
            [
                q(x["nome"]),
                x["anno"] if x["anno"].isdigit() else "null",
                q(x["data"]) if x["data"] else "null",
                q(x["oda"]),
                q(x["ente"]),
                q(x["corso"]),
                q(x["primo"]),
                q(x["attestato"]),
                q(x["data_att"]) if x["data_att"] else "null",
                q(x["scadenza"]) if x["scadenza"] else "null",
                q(x["dimessi"]),
                q(x["note"]),
            ]
        )
        + ")"
        for x in formazione_rows
    )

    sql = f"""-- Sync personale to 173 names from Excel + seed formazione from ELENCO TOT
with src_personale(full_name) as (
  values
{personale_values}
)
insert into public.personale (full_name, active)
select sp.full_name, true
from src_personale sp
where not exists (
  select 1
  from public.personale p
  where trim(regexp_replace(lower(p.full_name), '\\s+', ' ', 'g')) =
        trim(regexp_replace(lower(sp.full_name), '\\s+', ' ', 'g'))
);

with src(
  nome, anno, data, oda, ente, corso, primo_rilascio_aggiornamento, attestato,
  data_attestato, scadenza_attestato, dimessi, note
) as (
  values
{formazione_values}
), normalized as (
  select
    trim(regexp_replace(lower(nome), '\\s+', ' ', 'g')) as nome_norm,
    *
  from src
), personale_match as (
  select
    p.id_uuid as personale_id,
    trim(regexp_replace(lower(p.full_name), '\\s+', ' ', 'g')) as nome_norm
  from public.personale p
), mapped as (
  select
    pm.personale_id,
    n.anno::integer,
    n.data::date,
    nullif(n.oda, ''),
    nullif(n.ente, ''),
    n.corso,
    nullif(n.primo_rilascio_aggiornamento, ''),
    nullif(n.attestato, ''),
    n.data_attestato::date,
    n.scadenza_attestato::date,
    null::text as corso_prenotato,
    null::date as prima_data,
    null::date as seconda_data,
    null::text as orario,
    null::text as modalita,
    nullif(n.dimessi, ''),
    nullif(n.note, '')
  from normalized n
  join personale_match pm on pm.nome_norm = n.nome_norm
), dedup as (
  select distinct on (
    personale_id,
    corso,
    coalesce(data_attestato, '1900-01-01'::date),
    coalesce(scadenza_attestato, '1900-01-01'::date)
  )
    *
  from mapped
  order by
    personale_id,
    corso,
    coalesce(data_attestato, '1900-01-01'::date),
    coalesce(scadenza_attestato, '1900-01-01'::date),
    anno desc nulls last,
    data desc nulls last
)
insert into public.formazione_corsi (
  personale_id, anno, data, oda, ente, corso, primo_rilascio_aggiornamento, attestato,
  data_attestato, scadenza_attestato, corso_prenotato, prima_data, seconda_data, orario, modalita, dimessi, note
)
select * from dedup
on conflict (personale_id, corso, data_attestato) do update set
  anno = excluded.anno,
  data = excluded.data,
  oda = excluded.oda,
  ente = excluded.ente,
  primo_rilascio_aggiornamento = excluded.primo_rilascio_aggiornamento,
  attestato = excluded.attestato,
  scadenza_attestato = excluded.scadenza_attestato,
  dimessi = excluded.dimessi,
  note = excluded.note,
  updated_at = now();
"""

    OUT.write_text(sql, encoding="utf-8")
    print(f"personale names: {len(personale_names)}")
    print(f"formazione rows: {len(formazione_rows)}")
    print(f"written: {OUT}")


if __name__ == "__main__":
    main()


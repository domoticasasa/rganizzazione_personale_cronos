import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import requests
from openpyxl import load_workbook

SUPABASE_URL = "https://bjdimalvbdablzoctrsf.supabase.co"
ANON_KEY = (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImJqZGltYWx2YmRhYmx6b2N0cnNmIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzE5MzczNTcsImV4cCI6MjA4NzUxMzM1N30."
    "Lm4GHd1CWdSMMI87Wx2eBrHcmqFtMXO31hO7Q3BsYyc"
)
SRC = Path(r"c:\Users\alexandru.cibuc\Downloads\CORSI DI FORMAZIONE.xlsx")


def norm(text: str) -> str:
    t = (text or "").strip().lower().replace("’", "'")
    t = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")
    t = re.sub(r"[^a-z0-9 ]+", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def tkey(text: str) -> str:
    parts = [p for p in norm(text).split(" ") if p]
    parts.sort()
    return " ".join(parts)


def to_iso(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return ""
    s = str(v).strip()
    if not s:
        return ""
    try:
        return datetime.fromisoformat(s.replace("Z", "")).date().isoformat()
    except Exception:
        return ""


def s(v) -> str:
    return "" if v is None else str(v).strip()


def main() -> None:
    wb = load_workbook(SRC, data_only=True)
    ws_r = wb["RIEPILOGO"]
    ws_t = wb["ELENCO TOT"]

    riepilogo_names = {
        s(ws_r.cell(r, 4).value)
        for r in range(3, ws_r.max_row + 1)
        if s(ws_r.cell(r, 4).value)
    }
    riepilogo_keys = {tkey(n): n for n in riepilogo_names if tkey(n)}

    excel_rows = []
    for r in range(2, ws_t.max_row + 1):
        nome = s(ws_t.cell(r, 7).value)
        corso = s(ws_t.cell(r, 5).value)
        if not nome or not corso:
            continue
        excel_rows.append(
            {
                "name": nome,
                "key": tkey(nome),
                "corso": corso.strip().upper(),
                "ente": s(ws_t.cell(r, 4).value),
                "data_att": to_iso(ws_t.cell(r, 11).value),
                "scad": to_iso(ws_t.cell(r, 12).value),
            }
        )

    headers = {"apikey": ANON_KEY, "Authorization": f"Bearer {ANON_KEY}"}
    pr = requests.get(
        f"{SUPABASE_URL}/rest/v1/personale?select=id_uuid,full_name&limit=5000",
        headers=headers,
        timeout=30,
    )
    pr.raise_for_status()
    personale = pr.json()
    p_by_key = {tkey((p.get("full_name") or "")): p for p in personale if tkey((p.get("full_name") or ""))}

    fr = requests.get(
        f"{SUPABASE_URL}/rest/v1/formazione_corsi?select=personale_id,corso,ente,data_attestato,scadenza_attestato&limit=20000",
        headers=headers,
        timeout=30,
    )
    fr.raise_for_status()
    form = fr.json()

    id_to_key = {}
    for p in personale:
        k = tkey((p.get("full_name") or ""))
        if k and p.get("id_uuid"):
            id_to_key[p["id_uuid"]] = k

    db_index = defaultdict(list)
    for row in form:
        k = id_to_key.get(row.get("personale_id"))
        if not k:
            continue
        db_index[(k, (row.get("corso") or "").strip().upper())].append(row)

    missing_people = []
    for k, display in riepilogo_keys.items():
        if k not in p_by_key:
            missing_people.append(display)

    missing_courses = []
    missing_ente = []
    missing_data_att = []
    for x in excel_rows:
        key = (x["key"], x["corso"])
        rows = db_index.get(key, [])
        if not rows:
            missing_courses.append((x["name"], x["corso"]))
            continue
        if x["ente"] and not any((r.get("ente") or "").strip() for r in rows):
            missing_ente.append((x["name"], x["corso"]))
        if x["data_att"] and not any((r.get("data_attestato") or "").strip() for r in rows):
            missing_data_att.append((x["name"], x["corso"]))

    people_with_course_excel = {x["key"] for x in excel_rows if x["key"]}
    people_with_course_db = {id_to_key.get(r.get("personale_id")) for r in form if id_to_key.get(r.get("personale_id"))}
    no_course_in_db = sorted([riepilogo_keys[k] for k in riepilogo_keys if k not in people_with_course_db])

    print("RIEPILOGO names:", len(riepilogo_names))
    print("DB personale rows:", len(personale))
    print("DB formazione rows:", len(form))
    print("Missing people in DB personale:", len(missing_people))
    print("Missing course rows vs ELENCO TOT:", len(missing_courses))
    print("Missing ente where Excel has ente:", len(missing_ente))
    print("Missing data_attestato where Excel has it:", len(missing_data_att))
    print("People from RIEPILOGO without any course in DB:", len(no_course_in_db))
    print("Sample missing people:", missing_people[:20])
    print("Sample missing courses:", missing_courses[:20])
    print("Sample missing ente:", missing_ente[:20])
    print("Sample missing data_attestato:", missing_data_att[:20])
    print("No-course names sample:", no_course_in_db[:20])


if __name__ == "__main__":
    main()


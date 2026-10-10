from datetime import date, datetime
from pathlib import Path

import openpyxl


def norm(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).strip()
    if s in ("", "-", "--", "N/D"):
        return None
    return s


def num(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None


def sql(value):
    if value is None:
        return "NULL"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def main():
    repo = Path(r"c:\flutter_projects\organizzazione_personale_cronos - mod")
    xlsx = Path(r"C:\Users\alexandru.cibuc\Desktop\MEZI E DOT.xlsx")
    out = (
        repo
        / "supabase"
        / "migrations"
        / "20260422192500_seed_logistica_mdo_ferroviari_current_foglio1.sql"
    )

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = []
    for r in range(2, ws.max_row + 1):
        matricola = norm(ws.cell(r, 1).value)
        if not matricola:
            continue
        codice = norm(ws.cell(r, 2).value)
        descr_mezzo = norm(ws.cell(r, 3).value)
        descr_rumo = norm(ws.cell(r, 4).value)
        modello = norm(ws.cell(r, 5).value)
        equipment = norm(ws.cell(r, 6).value)
        matr_cost = norm(ws.cell(r, 7).value)
        cantiere_attuale = norm(ws.cell(r, 8).value)
        commessa = norm(ws.cell(r, 9).value)
        dispositivo_shuntaggio = norm(ws.cell(r, 69).value)
        lanterna_bilux = norm(ws.cell(r, 72).value)
        fanali_coda = norm(ws.cell(r, 75).value)
        tabella_coda = norm(ws.cell(r, 78).value)
        torcia_fiamma_rossa = norm(ws.cell(r, 81).value)
        bandiera_rossa_asta = norm(ws.cell(r, 84).value)
        scarpe_fermacarro = norm(ws.cell(r, 87).value)
        chiave_tripla_snodata = norm(ws.cell(r, 90).value)
        barra_traino = norm(ws.cell(r, 93).value)
        vaschetta_liquidi = norm(ws.cell(r, 96).value)
        dispositivo_shuntaggio_check = (
            str(dispositivo_shuntaggio).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if dispositivo_shuntaggio is not None
            else False
        )
        lanterna_bilux_check = (
            str(lanterna_bilux).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if lanterna_bilux is not None
            else False
        )
        fanali_coda_check = (
            str(fanali_coda).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if fanali_coda is not None
            else False
        )
        tabella_coda_check = (
            str(tabella_coda).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if tabella_coda is not None
            else False
        )
        torcia_fiamma_rossa_check = (
            str(torcia_fiamma_rossa).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if torcia_fiamma_rossa is not None
            else False
        )
        bandiera_rossa_asta_check = (
            str(bandiera_rossa_asta).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if bandiera_rossa_asta is not None
            else False
        )
        scarpe_fermacarro_check = (
            str(scarpe_fermacarro).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if scarpe_fermacarro is not None
            else False
        )
        chiave_tripla_snodata_check = (
            str(chiave_tripla_snodata).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if chiave_tripla_snodata is not None
            else False
        )
        barra_traino_check = (
            str(barra_traino).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if barra_traino is not None
            else False
        )
        vaschetta_raccolta_liquidi_check = (
            str(vaschetta_liquidi).strip().lower() in ("si", "sì", "yes", "ok", "x", "1", "true", "check", "ceck")
            if vaschetta_liquidi is not None
            else False
        )

        rows.append(
            (
                matricola,
                codice,
                descr_mezzo,
                descr_rumo,
                modello,
                equipment,
                matr_cost,
                cantiere_attuale,
                commessa,
                dispositivo_shuntaggio_check,
                lanterna_bilux_check,
                fanali_coda_check,
                tabella_coda_check,
                torcia_fiamma_rossa_check,
                bandiera_rossa_asta_check,
                scarpe_fermacarro_check,
                chiave_tripla_snodata_check,
                barra_traino_check,
                vaschetta_raccolta_liquidi_check,
            )
        )

    lines = []
    lines.append("delete from public.logistica_mdo_ferroviari;")
    lines.append("")
    lines.append("insert into public.logistica_mdo_ferroviari (")
    lines.append(
        "  matricola_interna, codice_identificativo_targa_rfi, descrizione_mezzo, descrizione_rumo, modello, equipment, matricola_costruttore, cantiere_attuale, commessa, dispositivo_shuntaggio_check, lanterna_bilux_check, fanali_coda_check, tabella_coda_check, torcia_fiamma_rossa_check, bandiera_rossa_asta_check, scarpe_fermacarro_check, chiave_tripla_snodata_check, barra_traino_check, vaschetta_raccolta_liquidi_check, active"
    )
    lines.append(") values")

    values_sql = []
    for row in rows:
        values_sql.append(
            "  ("
            + ", ".join([sql(v) for v in row] + ["true"])
            + ")"
        )
    lines.append(",\n".join(values_sql) + ";")
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {out} with {len(rows)} rows.")


if __name__ == "__main__":
    main()

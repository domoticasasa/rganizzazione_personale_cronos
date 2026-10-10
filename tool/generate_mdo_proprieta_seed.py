"""Genera righe SQL seed per logistica_mdo_proprieta da Excel."""
import openpyxl
from datetime import datetime
from pathlib import Path

EXCEL = Path(r"c:\Users\alexandru.cibuc\Desktop\logistica\MdO_Proprietà.xlsx")
OUT = Path(__file__).resolve().parent.parent / "supabase" / "_mdo_proprieta_seed_lines.txt"


def esc(s):
    if s is None:
        return "NULL"
    t = str(s).replace("'", "''").strip()
    if not t:
        return "NULL"
    return "'" + t + "'"


def esc_date(v):
    if v is None:
        return "NULL"
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d") + "'"
    s = str(v).strip()
    if not s:
        return "NULL"
    return esc(s)


def main():
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb["MDO"]
    gruppo = None
    ordine = 0
    rows_sql = []
    for r in range(2, ws.max_row + 1):
        codifica = ws.cell(r, 1).value
        if codifica is not None and str(codifica).strip():
            gruppo = str(codifica).strip()
            ordine = 0
        if not gruppo:
            continue
        ordine += 1
        tipologia = ws.cell(r, 2).value
        matricola = ws.cell(r, 3).value
        definizione = ws.cell(r, 4).value
        serie = ws.cell(r, 5).value
        dich = ws.cell(r, 6).value
        scad = ws.cell(r, 7).value
        ubic = ws.cell(r, 8).value
        is_princ = (
            "true"
            if (codifica is not None and str(codifica).strip())
            else "false"
        )
        cod = (
            esc(codifica)
            if (codifica is not None and str(codifica).strip())
            else "NULL"
        )
        rows_sql.append(
            f"  ({esc(gruppo)}, {cod}, {is_princ}, {ordine}, "
            f"{esc(tipologia)}, {esc(matricola)}, {esc(definizione)}, "
            f"{esc(serie)}, {esc(dich)}, {esc_date(scad)}, {esc(ubic)})"
        )

    OUT.write_text("\n".join(rows_sql), encoding="utf-8")
    print(f"Wrote {len(rows_sql)} rows to {OUT}")


if __name__ == "__main__":
    main()

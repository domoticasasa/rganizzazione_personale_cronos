import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\alexandru.cibuc\Desktop\Mod.DPIE_00.xlsx")
ws = wb[wb.sheetnames[0]]
for row in ws.iter_rows(min_row=1, max_row=120, min_col=1, max_col=30):
    for c in row:
        v = c.value
        if v is not None and str(v).strip() != "":
            print(f"{c.coordinate}\t{v}")

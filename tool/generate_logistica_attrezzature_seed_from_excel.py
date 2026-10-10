"""Genera migration SQL: svuota logistica_attrezzature e ricarica da Excel."""

from __future__ import annotations

import pathlib
from datetime import date, datetime

import openpyxl

EXCEL_PATH = r"c:\Users\alexandru.cibuc\Desktop\logistica\Attrezzature 2026.XLSX"
OUT_SQL = pathlib.Path(
    __file__).resolve().parent.parent / "supabase" / "migrations" / "20260523120000_reload_logistica_attrezzature_2026.sql"

# Fogli con colonne standard (esclusi duplicati / archivi / moduli legali).
STANDARD_SHEETS = (
    "Attrezzature 2026",
    "Cassetta Attrezzi",
)

# Fogli aggiuntivi mappati sullo schema attrezzature.
ALT_SHEETS = {
    "HILTI - CODICI ILLEGGIBILI": {
        "codice_cronos": "CODICE ARTICOLO",
        "descrizione_articolo": "DESCRIZIONE",
        "serial_number": "NUMERO SERIALE",
        "modello": "MODELLO",
    },
    "STR TARATURE": {
        "codice_cronos": "CODICE",
        "famiglia_attrezzi": "DESCRIZIONE BREVE",
        "serial_number": "SERIAL NUMBER",
        "descrizione_articolo": "DESCRIZIONE COMPLETA",
        "data_oda": "TARATURA",
        "stato_valore": "REVISIONIONE ANNUALE",
        "note": "NOTE",
    },
}


def norm_header(v: object) -> str:
    if v is None:
        return ""
    return str(v).strip().upper()


def norm_text(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    s = str(v).replace("\n", " ").strip()
    if s.lower() == "nan":
        return ""
    return " ".join(s.split())


def sql_text(v: object) -> str:
    s = norm_text(v)
    if not s:
        return "NULL"
    return "'" + s.replace("'", "''") + "'"


def is_effective_row(values: list[object]) -> bool:
    return any(norm_text(v) for v in values)


def row_values_standard(headers: list[str], raw: tuple[object, ...]) -> dict[str, object] | None:
    idx = {h: i for i, h in enumerate(headers) if h}

    def cell(name: str) -> object:
        i = idx.get(name)
        if i is None or i >= len(raw):
            return None
        return raw[i]

    codice = norm_text(cell("CODICE CRONOS"))
    serial = norm_text(cell("SERIAL NUMBER"))
    if not codice and not serial:
        return None

    return {
        "codice_cronos": cell("CODICE CRONOS"),
        "assegnatario": cell("ASSEGNATARIO"),
        "famiglia_attrezzi": cell("FAMIGLIA ATTREZZI"),
        "marca": cell("MARCA") if "MARCA" in idx else None,
        "descrizione_articolo": cell("DESCRIZIONE ARTICOLO"),
        "serial_number": cell("SERIAL NUMBER"),
        "modello": cell("MODELLO"),
        "oda": cell("ODA"),
        "data_oda": cell("DATA ODA"),
        "posizione": cell("POSIZIONE"),
        "ddt": cell("DDT"),
        "stato_valore": cell("STATO/VALORE"),
        "note": cell("NOTE"),
    }


def row_values_alt(headers: list[str], raw: tuple[object, ...], mapping: dict[str, str]) -> dict[str, object] | None:
    idx = {h: i for i, h in enumerate(headers) if h}

    def cell(hdr: str) -> object:
        i = idx.get(hdr)
        if i is None or i >= len(raw):
            return None
        return raw[i]

    out: dict[str, object] = {}
    for col, hdr in mapping.items():
        out[col] = cell(hdr)

    if not norm_text(out.get("codice_cronos")) and not norm_text(out.get("serial_number")):
        return None
    return out


def values_tuple(row: dict[str, object]) -> str:
    cols = [
        "codice_cronos",
        "assegnatario",
        "famiglia_attrezzi",
        "marca",
        "descrizione_articolo",
        "serial_number",
        "modello",
        "oda",
        "data_oda",
        "posizione",
        "ddt",
        "stato_valore",
        "note",
    ]
    return "(" + ", ".join(sql_text(row.get(c)) for c in cols) + ")"


def load_sheet(wb: openpyxl.Workbook, sheet_name: str) -> list[str]:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [norm_header(h) for h in rows[0]]
    values: list[str] = []

    if sheet_name in ALT_SHEETS:
        mapping = {k: norm_header(v) for k, v in ALT_SHEETS[sheet_name].items()}
        for raw in rows[1:]:
            vals = list(raw)
            if not is_effective_row(vals):
                continue
            parsed = row_values_alt(headers, tuple(vals), mapping)
            if parsed:
                values.append(values_tuple(parsed))
        return values

    required = [
        "CODICE CRONOS",
        "ASSEGNATARIO",
        "FAMIGLIA ATTREZZI",
        "DESCRIZIONE ARTICOLO",
        "SERIAL NUMBER",
        "MODELLO",
        "ODA",
        "DATA ODA",
        "POSIZIONE",
        "DDT",
        "STATO/VALORE",
        "NOTE",
    ]
    idx = {h: i for i, h in enumerate(headers) if h}
    missing = [h for h in required if h not in idx and h != "MARCA"]
    if "MARCA" not in idx and "MARCA" not in headers:
        missing.append("MARCA/Marca")
    if any(h for h in ["CODICE CRONOS", "DESCRIZIONE ARTICOLO"] if h not in idx):
        raise RuntimeError(f"Foglio {sheet_name}: intestazioni mancanti {missing}")

    for raw in rows[1:]:
        vals = list(raw)
        if not is_effective_row(vals):
            continue
        parsed = row_values_standard(headers, tuple(vals))
        if parsed:
            values.append(values_tuple(parsed))
    return values


def main() -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    all_values: list[str] = []
    stats: list[str] = []

    for sheet in STANDARD_SHEETS:
        if sheet not in wb.sheetnames:
            raise RuntimeError(f"Foglio mancante: {sheet}")
        chunk = load_sheet(wb, sheet)
        stats.append(f"-- {sheet}: {len(chunk)} righe")
        all_values.extend(chunk)

    for sheet in ALT_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        chunk = load_sheet(wb, sheet)
        stats.append(f"-- {sheet}: {len(chunk)} righe")
        all_values.extend(chunk)

    wb.close()

    if not all_values:
        raise RuntimeError("Nessun dato utile trovato nel file Excel.")

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    stats_block = "\n".join(stats)
    # Batch VALUES to avoid single huge statement limits.
    batch_size = 200
    insert_blocks: list[str] = []
    for i in range(0, len(all_values), batch_size):
        batch = all_values[i : i + batch_size]
        insert_blocks.append(
            "INSERT INTO public.logistica_attrezzature(\n"
            "  codice_cronos, assegnatario, famiglia_attrezzi, marca,\n"
            "  descrizione_articolo, serial_number, modello, oda, data_oda,\n"
            "  posizione, ddt, stato_valore, note, active\n"
            ")\nSELECT\n"
            "  v.codice_cronos, v.assegnatario, v.famiglia_attrezzi, v.marca,\n"
            "  v.descrizione_articolo, v.serial_number, v.modello, v.oda, v.data_oda,\n"
            "  v.posizione, v.ddt, v.stato_valore, v.note, true\n"
            "FROM (VALUES\n  "
            + ",\n  ".join(batch)
            + "\n) AS v(\n"
            "  codice_cronos, assegnatario, famiglia_attrezzi, marca,\n"
            "  descrizione_articolo, serial_number, modello, oda, data_oda,\n"
            "  posizione, ddt, stato_valore, note\n"
            ");"
        )

    sql = f"""-- Ricarica completa logistica_attrezzature da Attrezzature 2026.XLSX
-- Generato automaticamente il {generated_at}
{stats_block}
-- Totale righe: {len(all_values)}

DELETE FROM public.logistica_attrezzature;

{chr(10).join(insert_blocks)}
"""

    OUT_SQL.write_text(sql, encoding="utf-8")
    print(f"WROTE {OUT_SQL}")
    print(f"TOTAL_ROWS {len(all_values)}")
    for line in stats:
        print(line)


if __name__ == "__main__":
    main()
